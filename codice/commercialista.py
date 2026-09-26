"""
RIEPILOGO IN EXCEL E INVIO AL COMMERCIALISTA (funzione 7)
=======================================================================

   registro --> file XLSX del periodo --> email al commercialista
                                           con: il file allegato + i link PRIVATI agli originali su Drive

Il file Excel viene GENERATO DAL REGISTRO, con 3 fogli:
   - Riepilogo: totali del periodo (calcolati in centesimi, come sempre)
   - Fatture:   una riga per fattura attiva, con link all'originale su Drive
   - Annullate: le fatture annullate del periodo, con il motivo (trasparenza: lo storico non si nasconde)

L'indirizzo del commercialista NON è scritto nel codice: sta nel file .env
(destinatari parametrizzabili, non cablati nel codice).
"""

import os
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font

import config
import registro_fatture as reg
from google_servizi import ErroreGoogle, condividi_in_lettura, invia_email

CARTELLA_RIEPILOGHI = config.RIEPILOGHI
FORMATO_EURO = '#,##0.00 "€"'
NOMI_MESI = ["", "gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio",
             "agosto", "settembre", "ottobre", "novembre", "dicembre"]


def _nome_periodo(tipo: str, anno: int, numero: Optional[int]) -> str:
    if tipo == "mese":
        return f"{NOMI_MESI[numero]} {anno}"
    if tipo == "trimestre":
        return f"{numero}° trimestre {anno}"
    return f"anno {anno}"


def _euro(centesimi: Optional[int]):
    """Per Excel: da centesimi interi a euro ESATTI (Decimal, niente float)."""
    return None if centesimi is None else Decimal(centesimi) / 100


def _fatture_annullate_del_periodo(tipo: str, anno: int, numero: Optional[int]) -> list:
    risultato = []
    for f in reg.carica_registro():
        if not f["annullata"]:
            continue
        d = date.fromisoformat(f["data_emissione"])
        if d.year != anno:
            continue
        if tipo == "mese" and d.month != numero:
            continue
        if tipo == "trimestre" and (d.month - 1) // 3 + 1 != numero:
            continue
        risultato.append(f)
    return risultato


def esporta_riepilogo(tipo: str, anno: int, numero: Optional[int] = None) -> Path:
    """Crea il file Excel del periodo e restituisce il suo percorso."""
    if tipo not in ("mese", "trimestre", "anno"):
        raise ValueError("tipo deve essere 'mese', 'trimestre' o 'anno'")
    attive = sorted(reg.fatture_del_periodo(tipo, anno, numero), key=lambda f: f["data_emissione"])
    annullate = _fatture_annullate_del_periodo(tipo, anno, numero)
    grassetto = Font(bold=True)

    libro = Workbook()

    # --- Foglio 1: Riepilogo ---
    foglio = libro.active
    foglio.title = "Riepilogo"
    righe = [
        ("Riepilogo fatture fornitori", _nome_periodo(tipo, anno, numero)),
        ("Generato il", date.today().isoformat()),
        (),
        ("Fatture attive", len(attive)),
        ("Imponibile", _euro(sum(f["imponibile_cent"] or 0 for f in attive))),
        ("IVA", _euro(sum(f["iva_cent"] or 0 for f in attive))),
        ("Totale", _euro(sum(f["totale_cent"] or 0 for f in attive))),
        ("Di cui ancora da pagare", _euro(sum(f["totale_cent"] or 0 for f in attive if f["stato_pagamento"] == "da_pagare"))),
        ("Fatture annullate nel periodo", len(annullate)),
        ("Fatture con IVA non ordinaria", ", ".join(f"{f['id']} ({f['regime_iva'].replace('_', ' ')})"
                                                   for f in attive if f["regime_iva"] != "ordinario") or "nessuna"),
    ]
    for riga in righe:
        foglio.append(list(riga))
    for cella in ("B5", "B6", "B7", "B8"):
        foglio[cella].number_format = FORMATO_EURO
    foglio["A1"].font = grassetto
    foglio.column_dimensions["A"].width = 32
    foglio.column_dimensions["B"].width = 40

    # --- Foglio 2: Fatture ---
    intestazione = ["Id", "Fornitore", "P. IVA", "Numero", "Emissione", "Scadenza", "Imponibile",
                    "IVA", "Totale", "Regime IVA", "Stato", "Pagata il", "Originale su Drive"]
    foglio = libro.create_sheet("Fatture")
    foglio.append(intestazione)
    for f in attive:
        foglio.append([f["id"], f["fornitore"], f["partita_iva"], f["numero"], f["data_emissione"],
                       f["data_scadenza"], _euro(f["imponibile_cent"]), _euro(f["iva_cent"]),
                       _euro(f["totale_cent"]), f["regime_iva"].replace("_", " "),
                       f["stato_pagamento"].replace("_", " "), f["data_pagamento"], f.get("drive_link") or ""])
    for riga in foglio.iter_rows(min_row=2, min_col=7, max_col=9):
        for cella in riga:
            cella.number_format = FORMATO_EURO
    for cella in foglio[1]:
        cella.font = grassetto

    # --- Foglio 3: Annullate ---
    foglio = libro.create_sheet("Annullate")
    foglio.append(["Id", "Fornitore", "Numero", "Emissione", "Totale", "Motivo annullamento"])
    for f in annullate:
        foglio.append([f["id"], f["fornitore"], f["numero"], f["data_emissione"],
                       _euro(f["totale_cent"]), f["motivo_annullamento"]])
    for riga in foglio.iter_rows(min_row=2, min_col=5, max_col=5):
        for cella in riga:
            cella.number_format = FORMATO_EURO
    for cella in foglio[1]:
        cella.font = grassetto

    CARTELLA_RIEPILOGHI.mkdir(exist_ok=True)
    suffisso = f"{anno}" if tipo == "anno" else f"{anno}_{tipo}_{numero:02d}"
    percorso = CARTELLA_RIEPILOGHI / f"riepilogo_fatture_{suffisso}.xlsx"
    libro.save(percorso)
    return percorso


def invia_al_commercialista(tipo: str, anno: int, numero: Optional[int] = None) -> str:
    destinatario = os.getenv("COMMERCIALISTA_EMAIL", "").strip()
    if not destinatario:
        return "ERRORE: indirizzo del commercialista non configurato (COMMERCIALISTA_EMAIL nel file .env)."

    file_excel = esporta_riepilogo(tipo, anno, numero)
    attive = sorted(reg.fatture_del_periodo(tipo, anno, numero), key=lambda f: f["data_emissione"])
    periodo = _nome_periodo(tipo, anno, numero)

    # Link privati: ogni originale viene condiviso in sola lettura SOLO con il commercialista.
    righe_link, senza_originale = [], []
    for f in attive:
        if f.get("drive_id"):
            try:
                condividi_in_lettura(f["drive_id"], destinatario)
            except Exception:
                pass    # se la condivisione fallisce, il link resta comunque nell'email
            righe_link.append(f"- {f['fornitore']} n. {f['numero']} ({reg.centesimi_in_euro(f['totale_cent'])}): {f['drive_link']}")
        else:
            senza_originale.append(f"{f['id']} {f['fornitore']} n. {f['numero']}")

    testo = (f"Buongiorno,\n\nin allegato il riepilogo delle fatture fornitori del periodo: {periodo}.\n"
             f"Fatture attive: {len(attive)} - totale {reg.centesimi_in_euro(sum(f['totale_cent'] or 0 for f in attive))}.\n\n"
             "Documenti originali (link privati, accessibili solo a lei):\n"
             + ("\n".join(righe_link) if righe_link else "- nessun originale archiviato su Drive per questo periodo")
             + ("\n\nFatture senza originale su Drive: " + "; ".join(senza_originale) if senza_originale else "")
             + "\n\nCordiali saluti\n(messaggio generato dall'assistente fatture)")
    try:
        invia_email(destinatario, f"Riepilogo fatture fornitori - {periodo}", testo, allegato=file_excel)
    except ErroreGoogle as errore:
        return f"ERRORE: {errore} Il file è comunque pronto sul computer: {file_excel}"
    return (f"Email inviata a {destinatario} con il riepilogo {periodo} ({len(attive)} fatture, "
            f"{len(righe_link)} link a originali). File: {file_excel.name}")

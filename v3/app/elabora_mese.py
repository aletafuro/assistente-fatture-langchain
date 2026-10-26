"""
ELABORAZIONE DI UN MESE DI FATTURE (invio "a blocchi")
======================================================

L'amministrazione manda tutte le fatture del mese insieme (la prima settimana del mese dopo), più gli
estratti conto delle carte. Questo modulo:

  1. legge ogni PDF UNA volta sola (la lettura resta salvata in dati/letture/: rilanciare non costa);
  2. scarta le copie di cortesia quando c'è l'originale, e i documenti che non sono fatture
     (contabili di bonifico, avvisi, guide);
  3. per le fatture ESTERE calcola il controvalore in euro: carta Amex -> carta UniCredit -> cambio BdI;
  4. calcola la data di registrazione e la scadenza (dalla fattura o dai termini del fornitore);
  5. aggiorna l'anagrafica dei fornitori esteri e mette le modifiche nella colonna Note;
  6. scrive i file nel layout dell'amministrazione:
       - Registrazione contabile  (SOLO fatture estere)
       - Scadenziario fatture     (italiane ed estere)
       - copia dell'estratto Amex con importo in valuta, cambio, importo in euro e commissioni
       - un riepilogo di cosa va controllato a mano.

Il modello legge (punto 1); tutto il resto è Python.
Uso:  python app/elabora_mese.py CARTELLA [--amex FILE.xlsx] [--unicredit FILE.pdf] [--uscita CARTELLA]
"""

import argparse
import hashlib
import json
import re
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Optional

import anagrafica_esteri
import cambi_bdi
import config
import estratti_conto as ec
import termini
from estrazione import ErroreEstrazione, ServizioNonDisponibile, estrai_da_file
from fatture_comuni import PIVA_ESTERO_ASSENTE, e_estero

LETTURE = config.DATI / "letture"
CENT = Decimal("0.01")
# Parole che, nel nome di un file, indicano una contabile o un documento che NON è una fattura
# (in prova: "First instalment 4411 and full payment 4410.pdf", "Wire Transfer NWD 18.pdf").
NON_FATTURE = re.compile(r"\b(bonifico|contabile|guide|guida|avviso|estratto|statement|payment|instalment|"
                         r"installment|wire|transfer|remittance)\b", re.I)


# ---------------------------------------------------------------------------
# 1. LETTURA (con memoria: ogni file si legge una volta sola)
# ---------------------------------------------------------------------------
def _versione_lettura() -> str:
    """
    Cambia SOLO quando cambia quello che chiediamo al modello: istruzioni (che contengono anche il nome
    dell'azienda, dal .env) e schema dei campi. Si salva la sua risposta "grezza"; le conversioni Python
    (date, importi, partite IVA) si rifanno a ogni avvio. Così correggere Python non costa una nuova lettura.
    """
    import fatture_comuni
    domanda = config.MODELLO_TESTO + config.MODELLO_FOTO + fatture_comuni.ISTRUZIONI_ESTRAZIONE + json.dumps(fatture_comuni.DatiFattura.model_json_schema(),
                                                                sort_keys=True)
    return hashlib.sha256(domanda.encode()).hexdigest()[:8]


def _impronta(percorso: Path) -> str:
    return f"{hashlib.sha256(percorso.read_bytes()).hexdigest()[:20]}-{_versione_lettura()}"


def leggi(percorso: Path) -> dict:
    LETTURE.mkdir(parents=True, exist_ok=True)
    salvata = LETTURE / f"{_impronta(percorso)}.json"
    if salvata.exists():
        memoria = json.loads(salvata.read_text(encoding="utf-8"))
        if "errore" in memoria:
            return memoria | {"file": percorso.name}
        dati = estrai_da_file(percorso, grezzo=memoria["_grezzo"])      # niente chiamata al modello
    else:
        try:
            dati = estrai_da_file(percorso)
        except ServizioNonDisponibile as errore:          # temporaneo: niente salvataggio, si riprova
            return {"errore": str(errore), "file": percorso.name}
        except ErroreEstrazione as errore:
            dati = {"errore": str(errore)}
        salvata.write_text(json.dumps({"_grezzo": dati["_grezzo"]} if "_grezzo" in dati else dati,
                                      ensure_ascii=False, indent=1), encoding="utf-8")
    dati["file"] = percorso.name
    return dati


def file_da_leggere(cartella: Path, esclusi: set) -> list:
    """Tutti i PDF, senza le copie di cortesia che hanno anche l'originale e senza gli estratti conto."""
    pdf = [p for p in sorted(cartella.rglob("*")) if p.suffix.lower() in (".pdf", ".jpg", ".jpeg", ".png")
           and p.resolve() not in esclusi and "risultati" not in p.parts]
    originali = {re.sub(r"\s*\(copia di cortesia.*?\)", "", p.stem, flags=re.I).lower() for p in pdf
                 if "copia di cortesia" not in p.stem.lower()}
    scelti, saltati = [], []
    for p in pdf:
        base = re.sub(r"\s*\(copia di cortesia.*?\)", "", p.stem, flags=re.I).lower()
        if "copia di cortesia" in p.stem.lower() and base in originali:
            saltati.append((p, "copia di cortesia (c'è l'originale)"))
        elif NON_FATTURE.search(p.stem) and " ft n" not in p.stem.lower():
            saltati.append((p, "non è una fattura (dal nome del file)"))
        else:
            scelti.append(p)
    # Prima i file che si chiamano come fatture ("... ft n. ..."): se la stessa fattura compare anche in un
    # altro documento, vince la fattura.
    scelti.sort(key=lambda p: (" ft n" not in p.stem.lower() and "nota di credito" not in p.stem.lower(), str(p)))
    return scelti, saltati


# ---------------------------------------------------------------------------
# 2. REGOLE
# ---------------------------------------------------------------------------
def _euro(centesimi: Optional[int]) -> Optional[Decimal]:
    return None if centesimi is None else (Decimal(centesimi) / 100).quantize(CENT)


def carta_indicata(metodo: Optional[str]) -> Optional[str]:
    """Dal metodo di pagamento scritto sulla fattura: quale estratto conto guardare per primo."""
    m = (metodo or "").lower()
    if "american express" in m or "amex" in m:
        return "amex"
    if "mastercard" in m or "visa" in m:
        return "unicredit"
    return None


def numeri_compatibili(a: Optional[str], b: Optional[str]) -> bool:
    """Stesso numero di fattura, anche con zeri iniziali o separatori diversi ('0031415926' = '31415926').
    Se uno dei due manca, potrebbe essere la stessa fattura."""
    na, nb = (re.sub(r"[^0-9a-z]", "", (n or "").lower()).lstrip("0") for n in (a, b))
    return not na or not nb or na == nb


def pagamento_in_euro(fattura: dict, movimenti: list, usati: set):
    """Fattura in euro pagata con carta: addebito con lo stesso importo e lo stesso fornitore."""
    importo = _euro(fattura.get("totale_cent"))
    if importo is None or not fattura.get("data_emissione"):
        return None
    giorno = date.fromisoformat(fattura["data_emissione"])
    candidati = [m for m in movimenti if id(m) not in usati and m.importo_valuta is None and m.importo_eur == importo
                 and abs((m.data - giorno).days) <= 45
                 and ec.stesso_fornitore(fattura.get("fornitore") or "", m.descrizione)]
    if not candidati:
        return None
    migliore = min(candidati, key=lambda m: abs((m.data - giorno).days))
    usati.add(id(migliore))
    return migliore


def controvalore(fattura: dict, tutti: list, usati: set) -> tuple:
    """(importo in euro, fonte, movimento). Per le fatture in euro cerca solo se sono state pagate con carta."""
    importo = _euro(fattura.get("totale_cent"))
    movimenti = [m for m in tutti if m.importo_valuta is not None]
    if fattura.get("valuta") in (None, "EUR"):
        return importo, "in euro", pagamento_in_euro(fattura, tutti, usati)
    if importo is None or not fattura.get("data_emissione"):
        return None, "importo o data della fattura mancanti: controvalore da calcolare a mano", None
    ricerca = {"fornitore": fattura.get("fornitore") or "", "valuta": fattura["valuta"],
               "importo_valuta": importo, "data_emissione": fattura["data_emissione"]}
    preferita = carta_indicata(fattura.get("metodo_pagamento"))
    gruppi = [[m for m in movimenti if m.fonte == preferita], movimenti] if preferita else [movimenti]
    for gruppo in gruppi:
        trovato = ec.abbina(ricerca, gruppo, gia_usati=usati)
        if trovato:
            return trovato.controvalore_eur, f"{trovato.fonte} (cambio {trovato.cambio}, del {trovato.data:%d/%m})", trovato
    q = cambi_bdi.quotazione(fattura["valuta"], fattura["data_emissione"])
    if q is None:
        return None, "CAMBIO NON TROVATO", None
    return (importo / q[0]).quantize(CENT, ROUND_HALF_UP), f"cambio BdI {q[0]} del {date.fromisoformat(q[1]):%d/%m}", None


def nota_anagrafica(testo: str) -> str:
    """Nota breve per la colonna Note della registrazione (quello che serve all'amministrazione)."""
    nuovo = re.match(r"Creata la scheda anagrafica di .+ \(verifica partita IVA: (.+)\)\.", testo)
    if nuovo:
        return f"nuovo fornitore: creare l'anagrafica (P.IVA {nuovo.group(1)})"
    return testo.split(": ", 1)[-1].rstrip(".")       # "Rossi: nome cambiato da ..." -> "nome cambiato da ..."


def scadenza(fattura: dict) -> tuple:
    """(data di scadenza, da dove arriva)."""
    termine = termini.termine_di(fattura.get("fornitore") or "")
    scritta = fattura.get("data_scadenza")
    # Una "scadenza" uguale alla data fattura spesso vuol dire "non indicata" (es. fattura SDI con scadenza
    # alla data del documento): se il fornitore ha dei termini, valgono quelli (come fa l'amministrazione).
    if scritta and not (termine and scritta == fattura.get("data_emissione")):
        return scritta, "dalla fattura"
    if termine and fattura.get("data_emissione"):
        return termini.calcola_scadenza(fattura["data_emissione"], termine), f"termini {termine}"
    return None, "termini di pagamento da definire"


# ---------------------------------------------------------------------------
# 3. ELABORAZIONE DEL MESE
# ---------------------------------------------------------------------------
def elabora(cartella: Path, amex: Optional[Path] = None, unicredit: Optional[Path] = None) -> dict:
    esclusi = {p.resolve() for p in (amex, unicredit) if p}
    da_leggere, saltati = file_da_leggere(cartella, esclusi)
    movimenti = (ec.leggi_amex(amex) if amex else []) + (ec.leggi_unicredit(unicredit) if unicredit else [])
    in_valuta = [m for m in movimenti if m.importo_valuta is not None]
    usati, fatture, viste = set(), [], set()

    for percorso in da_leggere:
        f = leggi(percorso)
        f["percorso"] = str(percorso.relative_to(cartella))
        if f.get("errore"):
            saltati.append((percorso, f["errore"]))
            continue
        if f.get("tipo_documento") == "altro":
            saltati.append((percorso, "non è una fattura (secondo la lettura)"))
            continue
        chiave = (re.sub(r"\W", "", (f.get("partita_iva") or f.get("fornitore") or "").lower()),
                  re.sub(r"\W", "", (f.get("numero") or "").lower()))
        if chiave in viste:
            saltati.append((percorso, "stessa fattura già letta da un altro file"))
            continue
        # Stessa fattura con numero letto diverso (in prova: "... ft n. 1 - NEW.pdf" accanto all'originale):
        # stesso fornitore, stessa data e stesso importo -> un solo documento, si completano i dati mancanti.
        # ATTENZIONE: servono anche numeri compatibili. Due abbonamenti dello stesso giorno e importo
        # (Fabrikam, 2 x 15 USD il 21/09) sono due fatture diverse.
        gemella = next((g for g in fatture if g.get("data_emissione") == f.get("data_emissione")
                        and g.get("totale_cent") == f.get("totale_cent") and f.get("totale_cent")
                        and numeri_compatibili(g.get("numero"), f.get("numero"))
                        and ec.stesso_fornitore(g.get("fornitore") or "", f.get("fornitore") or "", alias={})), None)
        if gemella:
            for campo, valore in f.items():
                if gemella.get(campo) in (None, "") and valore not in (None, ""):
                    gemella[campo] = valore
            gemella["note"].append(f"documento presente due volte ({percorso.name})")
            saltati.append((percorso, f"doppione di {gemella['percorso']}"))
            continue
        viste.add(chiave)
        if not f.get("data_emissione") and f.get("fine_periodo"):
            f["data_emissione"] = f["fine_periodo"]
            f.setdefault("note_lettura", []).append("data fattura = fine del periodo di fatturazione")
        f["estera"] = e_estero(f.get("paese_fornitore")) and not f.get("formato_sdi")
        if f["estera"] and not f.get("partita_iva"):
            f["partita_iva"] = PIVA_ESTERO_ASSENTE
        segno = -1 if f.get("tipo_documento") == "nota_di_credito" else 1
        f["importo_valuta"] = _euro(f.get("totale_cent"))
        f["importo_eur"], f["fonte_cambio"], mov = controvalore(f, movimenti, usati)
        if mov:
            f["movimento"] = mov.per_json()
        if segno < 0:
            f["importo_valuta"] = -abs(f["importo_valuta"]) if f["importo_valuta"] is not None else None
            f["importo_eur"] = -abs(f["importo_eur"]) if f["importo_eur"] is not None else None
        if f.get("data_emissione"):
            f["data_registrazione"] = termini.data_registrazione(f["data_emissione"])
        f["scadenza"], f["fonte_scadenza"] = scadenza(f)
        if not f["scadenza"] and mov:                 # pagata con carta: la "scadenza" è il giorno dell'addebito
            f["scadenza"], f["fonte_scadenza"] = mov.data.isoformat(), "addebito su carta"
        # Solo per la scadenza scritta sulla fattura: un addebito su carta può venire prima (es. volo pagato
        # alla prenotazione), e non è un errore.
        if f["fonte_scadenza"] == "dalla fattura" and f.get("data_emissione") and f["scadenza"] < f["data_emissione"]:
            f.setdefault("note_lettura", []).append("scadenza precedente alla data fattura: controllare")
        f["note"] = list(f.get("note_lettura", []))
        if f["estera"]:
            try:
                cambiati = anagrafica_esteri.aggiorna_da_fattura(
                    f | {"id": f["file"], "data_emissione": f.get("data_emissione") or ""})
                f["note"] += [nota_anagrafica(n) for n in cambiati if "cambiato" in n or "Creata" in n]
            except Exception as errore:
                f["note"].append(f"anagrafica non aggiornata ({errore})")
        fatture.append(f)

    return {"fatture": fatture, "saltati": saltati, "movimenti": movimenti}


# ---------------------------------------------------------------------------
# 4. FILE IN USCITA (layout dell'amministrazione)
# ---------------------------------------------------------------------------
def _data_it(iso: Optional[str]):
    return date.fromisoformat(iso) if iso else None


def scrivi_file(risultato: dict, uscita: Path, mese: str, amex: Optional[Path] = None) -> list:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    uscita.mkdir(parents=True, exist_ok=True)
    creati = []

    def foglio(ws, intestazione, righe, formati=None):
        ws.append(intestazione)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in righe:
            ws.append(r)
        for i, colonna in enumerate(ws.columns):
            ws.column_dimensions[colonna[0].column_letter].width = min(55, max(12, *(len(str(c.value or "")) + 2 for c in colonna)))
            for c in colonna[1:]:
                if isinstance(c.value, date):
                    c.number_format = "DD/MM/YYYY"
                elif isinstance(c.value, (float, Decimal)):
                    c.number_format = "#,##0.00"

    # La pro-forma è una richiesta di pagamento: va nello scadenziario, NON nella registrazione contabile
    # (in prova: Proseware pro-forma 519-2026, messa dall'amministrazione nello scadenzario).
    estere = sorted([f for f in risultato["fatture"] if f["estera"] and f.get("tipo_documento") != "proforma"],
                    key=lambda f: (f.get("fornitore") or "").lower())
    wb = Workbook()
    foglio(wb.active, ["Data Registrazione", "Fornitore", "Numero Fattura", "Data Fattura",
                       "Importo Valuta Originaria", "Importo EURO", "Note"],
           [[_data_it(f.get("data_registrazione")), f.get("fornitore"), f.get("numero"), _data_it(f.get("data_emissione")),
             float(f["importo_valuta"]) if f["importo_valuta"] is not None else None,
             float(f["importo_eur"]) if f["importo_eur"] is not None else None,
             "; ".join(f["note"]) or None] for f in estere])
    wb.active.title = "Registrazione Contabile"
    creati.append(uscita / f"Registrazione contabile fatture estere {mese}.xlsx")
    wb.save(creati[-1])

    tutte = sorted(risultato["fatture"], key=lambda f: (f.get("scadenza") or "9999", (f.get("fornitore") or "").lower()))
    wb = Workbook()

    def nota_scadenziario(f):
        note = ["pro-forma: la fattura vera arriverà dopo"] if f.get("tipo_documento") == "proforma" else []
        if f.get("movimento"):
            m = f["movimento"]
            note.append(f"pagata con carta {m['fonte'].upper()} il {date.fromisoformat(m['data']):%d/%m/%Y}")
        if f["fonte_scadenza"] != "dalla fattura":
            note.append(f["fonte_scadenza"])
        return "; ".join(note) or None

    def riga(f):
        return [f.get("fornitore"), f.get("numero"), _data_it(f.get("data_emissione")),
                (float(f["importo_valuta"]) if f.get("valuta") not in (None, "EUR") and f["importo_valuta"] is not None else None),
                float(f["importo_eur"]) if f["importo_eur"] is not None else None,
                termini.termine_di(f.get("fornitore") or ""), _data_it(f.get("scadenza")), nota_scadenziario(f)]

    # Regola dell'amministrazione (call del 23/10): nello scadenziario vanno le fatture DA PAGARE (bonifico).
    # Quelle già addebitate su carta si pagano da sole: stanno in un foglio a parte, per consultarle.
    intestazione = ["Fornitore", "Numero Fattura", "Data Fattura", "Importo Valuta Originaria", "Importo EURO",
                    "Termini di Pagamento", "Data di Scadenza", "Note"]
    foglio(wb.active, intestazione, [riga(f) for f in tutte if not f.get("movimento")])
    wb.active.title = "Scadenziario fatture"
    foglio(wb.create_sheet("Pagate con carta"), intestazione, [riga(f) for f in tutte if f.get("movimento")])
    creati.append(uscita / f"Scadenziario fatture {mese}.xlsx")
    wb.save(creati[-1])

    if amex:
        # Copia dell'estratto Amex con le quattro colonne che l'amministrazione aggiungeva a mano.
        originale = load_workbook(amex, data_only=True).worksheets[0]
        righe = [list(r) for r in originale.iter_rows(values_only=True)]
        intest = next(i for i, r in enumerate(righe) if r and r[0] == "Data")
        col_f = righe[intest].index("Dettagli completi")
        wb = Workbook()
        ws = wb.active
        ws.title = "Dettagli transazione"
        for i, r in enumerate(righe):
            if i < intest:
                ws.append(r)
            elif i == intest:
                ws.append(r[:col_f + 1] + ["Importo valuta", "Valuta", "Cambio", "Importo EUR", "Commissioni"] + r[col_f + 1:])
            else:
                d = ec.leggi_dettagli_amex(r[col_f] or "") if r[col_f] else None
                extra = [None] * 5
                if d and d["importo_valuta"] and d["cambio"]:
                    eur = (d["importo_valuta"] / d["cambio"]).quantize(CENT, ROUND_HALF_UP)
                    extra = [float(d["importo_valuta"]), d["valuta"], float(d["cambio"]), float(eur),
                             float(Decimal(str(r[4])) - eur) if r[4] is not None else None]
                ws.append(r[:col_f + 1] + extra + r[col_f + 1:])
        creati.append(uscita / f"Estratto Amex con controvalori {mese}.xlsx")
        wb.save(creati[-1])

    righe = [f"Elaborazione {mese}: {len(risultato['fatture'])} fatture ({len(estere)} estere), "
             f"{len(risultato['saltati'])} file saltati.", "", "DA CONTROLLARE A MANO:"]
    for f in risultato["fatture"]:
        problemi = []
        if f["estera"] and f["importo_eur"] is None:
            problemi.append("controvalore in euro mancante")
        if not f.get("scadenza") and not f.get("movimento"):
            problemi.append("scadenza mancante (termini da definire, o pagata con una carta senza estratto conto)")
        if not f.get("numero") or not f.get("data_emissione"):
            problemi.append("numero o data mancanti")
        if problemi:
            righe.append(f"- {f['percorso']}: {', '.join(problemi)}")
    righe += ["", "FILE SALTATI:"] + [f"- {p.name}: {motivo}" for p, motivo in risultato["saltati"]]
    righe += ["", "CONTROVALORI DELLE FATTURE ESTERE:"] + [
        f"- {f.get('fornitore')} n. {f.get('numero')}: {f.get('importo_valuta')} {f.get('valuta')} -> "
        f"{f.get('importo_eur')} € ({f['fonte_cambio']})" for f in estere]
    creati.append(uscita / f"Riepilogo elaborazione {mese}.txt")
    creati[-1].write_text("\n".join(righe), encoding="utf-8")
    return creati


def main():
    parser = argparse.ArgumentParser(description="Elabora le fatture di un mese.")
    parser.add_argument("cartella", type=Path)
    parser.add_argument("--amex", type=Path)
    parser.add_argument("--unicredit", type=Path)
    parser.add_argument("--uscita", type=Path, default=config.RIEPILOGHI)
    parser.add_argument("--mese", default=date.today().strftime("%m-%Y"))
    a = parser.parse_args()
    risultato = elabora(a.cartella, a.amex, a.unicredit)
    for p in scrivi_file(risultato, a.uscita, a.mese, a.amex):
        print("creato:", p)


if __name__ == "__main__":
    main()

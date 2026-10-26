"""
RATEI E RISCONTI A UNA DATA DI CHIUSURA
=======================================

Calcolo A PARTE rispetto alla registrazione mensile (regola dell'amministrazione, 26/10): serve a fine anno
(31/12) per le scritture di chiusura, e a fine trimestre (31/03, 30/06, 30/09) o a una data qualsiasi per le
situazioni contabili intermedie.

Data di chiusura D, fattura con un periodo di competenza (letto dal modello nella descrizione):
   - RISCONTO: fattura datata entro D, con il periodo che va OLTRE D -> la quota dopo D si rinvia.
               Es. abbonamento annuale dal 17/09/2026 al 16/09/2027, chiusura 31/12/2026: 259 giorni su 365.
   - RATEO:    fattura datata DOPO D, con il periodo (in parte) entro D -> la quota fino a D è un costo del
               periodo chiuso (fatture da ricevere). Es. fattura di gennaio "dal 15/12/2026 al 15/01/2027",
               chiusura 31/12/2026: 17 giorni su 31.
La quota è proporzionale ai giorni ed è calcolata da Python, non dal modello. Estremi compresi, tranne quando
la fine ha lo stesso giorno del mese dell'inizio ("dal 15/12 al 15/01"): lì il periodo finisce il giorno prima.

Ogni blocco elaborato si salva in dati/blocchi/ (solo i campi che servono qui): così il calcolo si può rifare
in ogni momento con tutte le fatture arrivate fino a quel giorno. I ratei a una data sono completi solo
quando sono arrivate anche le fatture dei mesi dopo (per il 31/12: di solito entro il blocco di marzo).
"""

import json
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Optional

import config

BLOCCHI = config.DATI / "blocchi"
CAMPI = ("fornitore", "numero", "data_emissione", "inizio_periodo", "fine_periodo", "importo_eur", "valuta",
         "importo_valuta", "tipo_documento", "percorso")


# ---------------------------------------------------------------------------
# IL CALCOLO
# ---------------------------------------------------------------------------
def quota_alla_data(data_fattura: Optional[str], inizio: Optional[str], fine: Optional[str], importo,
                    chiusura: str) -> Optional[dict]:
    """Rateo o risconto di una fattura alla data di chiusura, oppure None se non ce n'è."""
    if not data_fattura or not (inizio or fine):
        return None
    inizio, fine = date.fromisoformat(inizio or fine), date.fromisoformat(fine or inizio)
    if fine < inizio:
        inizio, fine = fine, inizio
    # Periodi "dal 15/12 al 15/01" o "dal 15/11/2026 al 15/11/2027": il giorno finale è l'inizio del periodo
    # successivo, quindi non si conta (regola dell'amministrazione): 31 giorni, non 32; 365 (o 366), non 366.
    if fine.day == inizio.day and fine > inizio:
        fine -= timedelta(days=1)
    giorno_fattura, d = date.fromisoformat(data_fattura), date.fromisoformat(chiusura)
    if giorno_fattura <= d and fine > d:
        tipo, dal, al = "risconto", max(inizio, d + timedelta(days=1)), fine
    elif giorno_fattura > d and inizio <= d:
        tipo, dal, al = "rateo", inizio, min(fine, d)
    else:
        return None
    giorni, giorni_periodo = (al - dal).days + 1, (fine - inizio).days + 1
    quota = None if importo is None else \
        (Decimal(str(importo)) * giorni / giorni_periodo).quantize(Decimal("0.01"), ROUND_HALF_UP)
    return {"tipo": tipo, "dal": dal.isoformat(), "al": al.isoformat(), "giorni": giorni,
            "giorni_periodo": giorni_periodo, "quota": quota}


def chiusura_predefinita(oggi: Optional[date] = None) -> str:
    """L'ultima fine trimestre già passata (31/03, 30/06, 30/09, 31/12)."""
    oggi = oggi or date.today()
    fine_trimestri = [date(oggi.year - 1, 12, 31)] + [date(oggi.year, m, g) for m, g in ((3, 31), (6, 30), (9, 30), (12, 31))]
    return max(d for d in fine_trimestri if d < oggi).isoformat()


# ---------------------------------------------------------------------------
# I BLOCCHI SALVATI
# ---------------------------------------------------------------------------
def salva_blocco(mese: str, fatture: list) -> Path:
    """mese 'MM-AAAA'. Rielaborare lo stesso mese sostituisce il file (niente doppioni)."""
    BLOCCHI.mkdir(parents=True, exist_ok=True)
    percorso = BLOCCHI / f"{mese[3:]}-{mese[:2]}.json"           # 2026-09.json: in ordine di mese
    righe = [{c: (str(f[c]) if isinstance(f.get(c), Decimal) else f.get(c)) for c in CAMPI} for f in fatture]
    percorso.write_text(json.dumps({"mese": mese, "fatture": righe}, ensure_ascii=False, indent=1), encoding="utf-8")
    return percorso


def calcola(chiusura: str) -> dict:
    """Ratei e risconti alla data, con tutte le fatture dei blocchi salvati."""
    righe, blocchi = [], sorted(BLOCCHI.glob("*.json")) if BLOCCHI.exists() else []
    for percorso in blocchi:
        for f in json.loads(percorso.read_text(encoding="utf-8"))["fatture"]:
            q = quota_alla_data(f.get("data_emissione"), f.get("inizio_periodo"), f.get("fine_periodo"),
                                f.get("importo_eur"), chiusura)
            if q:
                righe.append(f | q)
    righe.sort(key=lambda r: (r["tipo"], (r.get("fornitore") or "").lower()))
    return {"chiusura": chiusura, "righe": righe, "blocchi": [p.stem for p in blocchi]}


def nota(r: dict) -> str:
    """'risconto: 4.630,14 € (260/365 giorni, dal 01/01/2027 al 17/09/2027)'."""
    importo = "importo da calcolare" if r["quota"] is None else \
        f"{Decimal(str(r['quota'])):,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")
    return (f"{r['tipo']}: {importo} ({r['giorni']}/{r['giorni_periodo']} giorni, "
            f"dal {date.fromisoformat(r['dal']):%d/%m/%Y} al {date.fromisoformat(r['al']):%d/%m/%Y})")


def scrivi_excel(risultato: dict, cartella: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    cartella.mkdir(parents=True, exist_ok=True)
    d = date.fromisoformat(risultato["chiusura"])
    wb = Workbook()
    ws = wb.active
    ws.title = f"Al {d:%d-%m-%Y}"
    ws.append(["Tipo", "Fornitore", "Numero Fattura", "Data Fattura", "Inizio periodo", "Fine periodo",
               "Importo EURO", "Dal", "Al", "Giorni", "Giorni del periodo", "Quota EURO"])
    for c in ws[1]:
        c.font = Font(bold=True)
    giorno = lambda iso: date.fromisoformat(iso) if iso else None                      # noqa: E731
    for r in risultato["righe"]:
        ws.append([r["tipo"], r.get("fornitore"), r.get("numero"), giorno(r.get("data_emissione")),
                   giorno(r.get("inizio_periodo")), giorno(r.get("fine_periodo")),
                   float(r["importo_eur"]) if r.get("importo_eur") is not None else None,
                   giorno(r["dal"]), giorno(r["al"]), r["giorni"], r["giorni_periodo"],
                   float(r["quota"]) if r["quota"] is not None else None])
    for colonna, larghezza in zip("ABCDEFGHIJKL", (10, 32, 20, 12, 12, 12, 13, 12, 12, 8, 10, 12)):
        ws.column_dimensions[colonna].width = larghezza
    for riga in ws.iter_rows(min_row=2):
        for c in riga:
            if isinstance(c.value, date):
                c.number_format = "DD/MM/YYYY"
            elif isinstance(c.value, float):
                c.number_format = "#,##0.00"
    totali = {t: sum(float(r["quota"] or 0) for r in risultato["righe"] if r["tipo"] == t) for t in ("rateo", "risconto")}
    ws.append([])
    ws.append(["Totale ratei", None, None, None, None, None, None, None, None, None, None, totali["rateo"]])
    ws.append(["Totale risconti", None, None, None, None, None, None, None, None, None, None, totali["risconto"]])
    ws.append([f"Calcolati con i blocchi: {', '.join(risultato['blocchi']) or 'nessuno'}"])
    percorso = cartella / f"Ratei e risconti al {d:%d-%m-%Y}.xlsx"
    wb.save(percorso)
    return percorso


def messaggio(risultato: dict) -> str:
    d = date.fromisoformat(risultato["chiusura"])
    ratei = [r for r in risultato["righe"] if r["tipo"] == "rateo"]
    risconti = [r for r in risultato["righe"] if r["tipo"] == "risconto"]
    testo = [f"Ratei e risconti al {d:%d/%m/%Y}: {len(ratei)} ratei e {len(risconti)} risconti "
             f"(blocchi usati: {', '.join(risultato['blocchi']) or 'nessuno'})."]
    testo += [f"- {r.get('fornitore')} n. {r.get('numero')}: {nota(r)}" for r in risultato["righe"][:30]]
    if not ratei:
        testo.append("Per i ratei servono anche le fatture arrivate dopo la data: rifai il calcolo dopo i blocchi "
                     "dei mesi successivi.")
    return "\n".join(testo)

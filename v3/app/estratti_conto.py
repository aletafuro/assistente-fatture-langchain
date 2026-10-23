"""
ESTRATTI CONTO DELLE CARTE: controvalore in euro delle fatture estere pagate con carta
====================================================================================

Regola dell'amministrazione per il controvalore in euro di una fattura in valuta:
   1. pagata con AMERICAN EXPRESS -> cambio applicato da Amex (estratto conto Excel);
   2. pagata con carta UNICREDIT  -> importo in euro scritto sull'estratto conto (PDF);
   3. pagata con BONIFICO          -> cambio Banca d'Italia del giorno della fattura (cambi_bdi.py).

Amex: nella colonna "Dettagli completi" c'è una sola stringa, per esempio
   "Foreign Spend Amount: 99.00 Dollari Statunitensi Commission Amount: 2.18 Currency Exchange Rate: 1.1332"
   importo in euro = 99,00 / 1,1332 = 87,36   (le commissioni sono a parte: addebito 89,54 - 87,36 = 2,18)
   Le parti possono cambiare posizione e lunghezza: ogni dato si cerca con la sua etichetta.

UniCredit: ogni operazione ha data registrazione, riferimento, descrizione, data operazione, importo in valuta,
   divisa, cambio e importo in euro; la commissione di cambio è una riga separata.

Abbinamento fattura <-> movimento: STESSO IMPORTO IN VALUTA (al centesimo), poi nome del fornitore
(anche tramite alias: sulla carta "CODEWAVE", la fattura è di "Fabrikam") e data più vicina.
Se più movimenti hanno lo stesso importo (es. abbonamenti OpenAI), vince la data più vicina.
"""

import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Optional

import config

ALIAS = config.DATI / "alias_fornitori.json"

# Amex scrive la valuta per esteso in italiano ("Dollari Statunitensi", "Corone Cecoslovacche"...):
# si riconosce dalle parole, non dalla frase esatta.
VALUTE_AMEX = [("statunitens", "USD"), ("canades", "CAD"), ("australian", "AUD"), ("sterlin", "GBP"),
               ("svizzer", "CHF"), ("yen", "JPY"), ("ceca", "CZK"), ("cech", "CZK"), ("cecoslovacc", "CZK"),
               ("zlot", "PLN"), ("svedes", "SEK"), ("danes", "DKK"), ("norveges", "NOK"), ("fiorin", "HUF"),
               ("ungheres", "HUF"), ("rupi", "INR"), ("real", "BRL"), ("pesos messican", "MXN"), ("shekel", "ILS")]


def valuta_amex(nome: str) -> Optional[str]:
    nome = (nome or "").lower()
    if re.fullmatch(r"[a-z]{3}", nome):
        return nome.upper()
    return next((codice for parola, codice in VALUTE_AMEX if parola in nome), None)


@dataclass
class Movimento:
    fonte: str                       # "amex" / "unicredit"
    data: date                       # data dell'operazione
    descrizione: str
    importo_eur: Decimal             # addebito in carta (Amex: comprese le commissioni)
    importo_valuta: Optional[Decimal] = None
    valuta: Optional[str] = None
    cambio: Optional[Decimal] = None
    controvalore_eur: Optional[Decimal] = None   # euro della fattura, SENZA commissioni
    commissioni: Optional[Decimal] = None
    riferimento: Optional[str] = None
    extra: dict = field(default_factory=dict)

    def per_json(self) -> dict:
        d = asdict(self)
        return {k: (str(v) if isinstance(v, (Decimal, date)) else v) for k, v in d.items()}


def _numero(testo) -> Optional[Decimal]:
    if testo in (None, ""):
        return None
    if isinstance(testo, (int, float, Decimal)):
        return Decimal(str(testo))
    t = str(testo).strip().replace(" ", "")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    return Decimal(t)


def _due_decimali(x: Decimal) -> Decimal:
    return x.quantize(Decimal("0.01"), ROUND_HALF_UP)


# ---------------------------------------------------------------------------
# AMERICAN EXPRESS (Excel)
# ---------------------------------------------------------------------------
_AMEX = {
    "importo": re.compile(r"Foreign Spend Amount:\s*([\d.,]+)\s*(.*?)\s*(?=Commission Amount|Currency Exchange Rate|$)", re.I),
    "commissioni": re.compile(r"Commission Amount:\s*([\d.,]+)", re.I),
    "cambio": re.compile(r"Currency Exchange Rate:\s*([\d.,]+)", re.I),
}


def leggi_dettagli_amex(testo: str) -> Optional[dict]:
    """Dalla stringa della colonna 'Dettagli completi' a importo in valuta, valuta, commissioni e cambio."""
    if not testo or "Foreign Spend Amount" not in testo:
        return None
    importo = _AMEX["importo"].search(testo)
    cambio = _AMEX["cambio"].search(testo)
    commissioni = _AMEX["commissioni"].search(testo)
    nome_valuta = importo.group(2).strip().lower() if importo else ""
    return {"importo_valuta": _numero(importo.group(1)) if importo else None,
            "valuta": valuta_amex(nome_valuta),
            "commissioni": _numero(commissioni.group(1)) if commissioni else None,
            "cambio": _numero(cambio.group(1)) if cambio else None}


def leggi_amex(percorso: Path, foglio: Optional[str] = None) -> list:
    """Movimenti dell'estratto conto Amex in Excel (il foglio originale scaricato dal portale Amex)."""
    from openpyxl import load_workbook
    libro = load_workbook(percorso, read_only=True, data_only=True)
    ws = libro[foglio] if foglio else libro.worksheets[0]
    righe = list(ws.iter_rows(values_only=True))
    intestazione = next(i for i, r in enumerate(righe) if r and r[0] == "Data")
    colonne = {nome: i for i, nome in enumerate(righe[intestazione]) if nome}
    movimenti = []
    for r in righe[intestazione + 1:]:
        if not r or not r[colonne["Data"]] or r[colonne["Importo"]] in (None, ""):
            continue
        quando = r[colonne["Data"]]
        quando = quando.date() if isinstance(quando, datetime) else datetime.strptime(str(quando), "%m/%d/%Y").date()
        m = Movimento("amex", quando, " ".join(str(r[colonne["Descrizione"]]).split()), _numero(r[colonne["Importo"]]),
                      riferimento=r[colonne.get("Riferimento")] if "Riferimento" in colonne else None,
                      extra={"titolare": r[colonne["Titolare"]] if "Titolare" in colonne else None,
                             "paese": r[colonne["Paese"]] if "Paese" in colonne else None})
        dettagli = leggi_dettagli_amex(r[colonne["Dettagli completi"]] if "Dettagli completi" in colonne else "")
        if dettagli and dettagli["importo_valuta"] and dettagli["cambio"]:
            m.importo_valuta, m.valuta, m.cambio = dettagli["importo_valuta"], dettagli["valuta"], dettagli["cambio"]
            m.controvalore_eur = _due_decimali(m.importo_valuta / m.cambio)
            m.commissioni = m.importo_eur - m.controvalore_eur      # come fa l'amministrazione: addebito - controvalore
            m.extra["commissioni_amex"] = str(dettagli["commissioni"])
        movimenti.append(m)
    return movimenti


# ---------------------------------------------------------------------------
# UNICREDIT (PDF)
# ---------------------------------------------------------------------------
_UNICREDIT = re.compile(
    r"(?P<reg>\d{2}\.\d{2}\.\d{2}) (?P<rif>H\d{5}) (?P<descr>.+?) (?P<op>\d{2}\.\d{2}\.\d{2}) "
    r"(?P<importo>[\d.]+,\d{2})(?: (?P<valuta>[A-Z]{3}) (?P<cambio>[\d.]+,\d+) (?P<eur>[\d.]+,\d{2}))?(?=\s|$)")


def leggi_unicredit(percorso: Path) -> list:
    """Movimenti dell'estratto conto UniCredit (PDF). Le righe 'Commissione per cambio' finiscono sul movimento."""
    from fatture_comuni import leggi_pdf
    testo = " ".join(leggi_pdf(Path(percorso)).split())
    movimenti, per_riferimento = [], {}
    for m in _UNICREDIT.finditer(testo):
        quando = datetime.strptime(m["op"], "%d.%m.%y").date()
        descr = m["descr"].strip()
        if descr.lower().startswith("commissione per cambio"):
            if m["rif"] in per_riferimento:
                per_riferimento[m["rif"]].commissioni = _numero(m["importo"])
            continue
        if m["valuta"]:
            mov = Movimento("unicredit", quando, descr, _numero(m["eur"]), _numero(m["importo"]), m["valuta"],
                            _numero(m["cambio"]), _numero(m["eur"]), riferimento=m["rif"])
        else:
            mov = Movimento("unicredit", quando, descr, _numero(m["importo"]), riferimento=m["rif"])
        movimenti.append(mov)
        per_riferimento[m["rif"]] = mov
    return movimenti


# ---------------------------------------------------------------------------
# ALIAS e ABBINAMENTO
# ---------------------------------------------------------------------------
def _parole(testo: str) -> set:
    vuote = {"inc", "llc", "ltd", "limited", "gmbh", "srl", "spa", "bv", "nv", "the", "com", "io", "ai", "sa",
             "san", "francisco", "new", "york", "dublin", "amsterdam", "subscr", "subscription", "uk", "ft"}
    return {p for p in re.findall(r"[a-z0-9]{3,}", (testo or "").lower()) if p not in vuote}


def carica_alias() -> dict:
    return json.loads(ALIAS.read_text(encoding="utf-8")) if ALIAS.exists() else {}


def aggiungi_alias(sulla_carta: str, fornitore: str) -> str:
    """Es. aggiungi_alias('codewave', 'Fabrikam'): da ora 'CODEWAVE' sull'estratto conto = Fabrikam."""
    alias = carica_alias()
    alias[sulla_carta.strip().lower()] = fornitore.strip()
    ALIAS.write_text(json.dumps(alias, ensure_ascii=False, indent=2), encoding="utf-8")
    return f"Ricordato: \"{sulla_carta}\" sull'estratto conto corrisponde al fornitore {fornitore}."


def _simili(a: set, b: set) -> bool:
    """Almeno una parola in comune, tollerando piccoli errori di battitura ("Fabrikm" = "Fabrikam")."""
    from difflib import SequenceMatcher
    return any(x == y or (min(len(x), len(y)) >= 5 and SequenceMatcher(None, x, y).ratio() >= 0.85)
               for x in a for y in b)


def stesso_fornitore(fornitore: str, descrizione: str, alias: Optional[dict] = None) -> bool:
    alias = carica_alias() if alias is None else alias
    descr = (descrizione or "").lower()
    for sulla_carta, nome in alias.items():
        if sulla_carta in descr and _simili(_parole(nome), _parole(fornitore)):
            return True
    return _simili(_parole(fornitore), _parole(descrizione))


def abbina(fattura: dict, movimenti: list, alias: Optional[dict] = None, gia_usati: Optional[set] = None) -> Optional[Movimento]:
    """
    Movimento di carta che paga questa fattura in valuta, oppure None (-> cambio Banca d'Italia).
    fattura: valuta, importo in valuta (decimale, es. Decimal('99.00')), data_emissione, fornitore.
    """
    gia_usati = gia_usati if gia_usati is not None else set()
    importo = fattura["importo_valuta"]
    giorno = date.fromisoformat(fattura["data_emissione"])
    candidati = [m for m in movimenti if id(m) not in gia_usati and m.importo_valuta is not None
                 and m.valuta == fattura["valuta"] and m.importo_valuta == importo]
    if not candidati:
        # Valuta letta male (in prova: "CSK" sulla fattura, letto come GBP): stesso importo e stesso fornitore
        # bastano, e la valuta giusta è quella del movimento.
        candidati = [m for m in movimenti if id(m) not in gia_usati and m.importo_valuta == importo
                     and stesso_fornitore(fattura["fornitore"], m.descrizione, alias)]
    if not candidati:
        return None
    nome_ok = [m for m in candidati if stesso_fornitore(fattura["fornitore"], m.descrizione, alias)]
    scelti = nome_ok or (candidati if len(candidati) == 1 else [])
    if not scelti:
        return None
    migliore = min(scelti, key=lambda m: abs((m.data - giorno).days))
    if abs((migliore.data - giorno).days) > 45:          # troppo lontano: non è il pagamento di questa fattura
        return None
    gia_usati.add(id(migliore))
    return migliore

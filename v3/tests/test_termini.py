"""Data di registrazione, scadenze dai termini e dal formato SDI, rate (regole dell'amministrazione)."""

import json
from datetime import date

import pytest

import termini
from fatture_comuni import formato_sdi, scadenze_sdi


@pytest.mark.parametrize("riferimento, registrazione", [
    ("2026-09-15", "2026-09-30"),     # 30/09/2026 è mercoledì
    ("2027-01-10", "2027-01-29"),     # esempio dell'amministrazione: 31 domenica, 30 sabato -> venerdì 29
    ("2026-10-02", "2026-10-30"),     # 31/10/2026 è sabato -> venerdì 30
    ("2026-02-03", "2026-02-27"),     # 28/02/2026 è sabato
    ("10-2026", "2026-10-30"),        # mese del blocco, scritto come nei nomi dei file
])
def test_data_registrazione(riferimento, registrazione):
    assert termini.data_registrazione(riferimento) == registrazione


def test_festivi_rossi_del_calendario():
    # Pasqua 2027 = 28 marzo, Pasquetta = 29 marzo; Ferragosto; Natale e Santo Stefano.
    assert termini._pasqua(2026) == date(2026, 4, 5) and termini._pasqua(2027) == date(2027, 3, 28)
    assert not termini.lavorativo(date(2027, 3, 29))
    assert not termini.lavorativo(date(2025, 8, 15)) and termini.lavorativo(date(2025, 8, 14))
    # Dicembre 2027: il 31 è venerdì ma lavorativo; dicembre 2022: 31 sabato -> 30 venerdì.
    assert termini.data_registrazione("12-2027") == "2027-12-31"
    assert termini.data_registrazione("12-2022") == "2022-12-30"


@pytest.mark.parametrize("termine, attesa", [
    ("30 gg DF", (30, False)), ("60 gg DF", (60, False)), ("BONIFICO B. 30 GG. FM.", (30, True)),
    ("60 gg D.F.F.M.", (60, True)), ("a vista", None),
])
def test_lettura_termini(termine, attesa):
    assert termini.leggi_termine(termine) == attesa


def test_calcolo_scadenza():
    assert termini.calcola_scadenza("2026-09-30", "30 gg DF") == "2026-10-30"
    assert termini.calcola_scadenza("2026-09-15", "30 gg FM") == "2026-10-31"


def test_termini_per_fornitore_con_nome_diverso():
    termini.imposta("Fabrikam", "30 gg DF")
    assert termini.termine_di("Fabrikam Srl") == "30 gg DF"
    assert termini.termine_di("Sconosciuto SpA") is None
    assert "Non capisco" in termini.imposta("Rossi", "presto")


def test_termini_dal_file_excel_dell_amministrazione(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook
    percorso = tmp_path / "Termini.xlsx"
    wb = Workbook()
    wb.active.append(["Fornitore", "Numero Fattura", "Data Fattura", "Importo", "Data di Scadenza"])
    wb.active.append(["Northwind BV", None, None, None, "60 gg DF"])
    wb.save(percorso)
    monkeypatch.setattr(termini, "FILE_TERMINI", percorso)
    assert termini.termine_di("Northwind B.V.") == "60 gg DF"        # letto dal file, senza importarlo
    # Detto in chat: finisce anche nel file, nella stessa colonna che usa l'amministrazione.
    assert "anche nel file" in termini.imposta("Contoso Ltd", "14 gg DF")
    righe = list(load_workbook(percorso).active.iter_rows(values_only=True))
    assert righe[-1][0] == "Contoso Ltd" and righe[-1][4] == "14 gg DF"


def test_rate_dette_a_parole_e_regola_del_fornitore(monkeypatch, tmp_path):
    regole = tmp_path / "regole_scadenze.json"
    regole.write_text(json.dumps([
        {"fornitore": "Contoso Media", "se_contiene": "chat",
         "rate": [{"percentuale": 60, "termine": "7 gg DF"}, {"percentuale": 40, "termine": "30 gg FM"}]},
        {"fornitore": "Contoso Media", "se_contiene": "mail", "rate": [{"percentuale": 100, "termine": "7 gg DF"}]}]))
    monkeypatch.setattr(termini, "REGOLE", regole)
    assert termini.leggi_rate("60% a 7 gg DF e 40% a 30 gg FM") == [
        {"percentuale": 60, "termine": "7 gg DF"}, {"percentuale": 40, "termine": "30 gg FM"}]
    assert termini.leggi_rate("50% a 30 gg DF") is None                 # non fa 100: non sono rate
    assert termini.regola_per("Contoso Media Srl", "Servizio CHAT settembre")[0]["percentuale"] == 60
    assert termini.regola_per("Contoso Media Srl", "Servizio mail")[0]["percentuale"] == 100
    assert termini.regola_per("Altro fornitore", "chat") is None
    piano = termini.piano_di_pagamento("2026-09-30", 100001, termini.regola_per("Contoso Media", "chat"))
    assert [(r["scadenza"], r["importo_cent"]) for r in piano] == [("2026-10-07", 60001), ("2026-10-31", 40000)]


SDI = """Cedente/prestatore (fornitore)
Identificativo fiscale ai fini IVA: IT03792180980
Cessionario/committente (cliente)
Modalità pagamento Dettagli Scadenze Importo
MP05 Bonifico IBAN IT97O0503455340000000007075
Data termine 15-09-2026 46gg"""


def test_formato_sdi_e_scadenza_con_giorni():
    assert formato_sdi(SDI)
    assert not formato_sdi("Invoice number 123 Bill to")
    assert scadenze_sdi(SDI) == ["2026-10-31"]          # 15/09 + 46 giorni
    assert scadenze_sdi("Data scadenza 30-11-2026") == ["2026-11-30"]


def test_manca_la_nostra_partita_iva():
    from fatture_comuni import piva_cliente_presente
    assert piva_cliente_presente("Bill to Studio Prova, VAT IT 012 345 678 97") is True    # azienda dei test
    assert piva_cliente_presente("Bill to Studio Prova, Milano") is False
    assert piva_cliente_presente(None) is None                                         # foto: non si sa


def test_rateo_fattura_di_gennaio_con_competenza_dicembre():
    # Esempio dell'amministrazione: fattura di gennaio 2027 "dal 15/12/2026 al 15/01/2027".
    [q] = termini.quote_fuori_anno("2027-01-20", "2026-12-15", "2027-01-15", "320.00")
    assert (q["tipo"], q["anno"], q["giorni"], q["giorni_periodo"], str(q["quota"])) == ("rateo", 2026, 17, 32, "170.00")
    assert termini.nota_competenza(q) == \
        "rateo: 170,00 € di competenza 2026 (17/32 giorni, dal 15/12/2026 al 31/12/2026)"
    # Fattura di febbraio "competenza dicembre 2026": tutta dell'anno prima.
    [q] = termini.quote_fuori_anno("2027-02-05", "2026-12-01", "2026-12-31", "1000")
    assert q["giorni"] == q["giorni_periodo"] == 31 and str(q["quota"]) == "1000.00"


def test_risconto_e_periodo_tutto_nell_anno():
    [q] = termini.quote_fuori_anno("2026-12-10", "2026-12-10", "2027-01-09", "310")
    assert (q["tipo"], q["anno"], q["giorni"], str(q["quota"])) == ("risconto", 2027, 9, "90.00")
    assert termini.quote_fuori_anno("2026-09-30", "2026-09-01", "2026-09-30", "100") == []
    assert termini.quote_fuori_anno("2026-09-30", None, None, "100") == []
